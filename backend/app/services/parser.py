import csv
import io
import re
import zipfile
from collections.abc import Iterable
from pathlib import Path

import pdfplumber
import pypdfium2 as pdfium
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import get_settings
from app.domain.document_ir import (
    DocumentNeedsVision,
    ParsedBlock,
    ParsedCell,
    ParsedDocument,
    ParsedPage,
)
from app.services.office_preview import render_office_pdf
from app.services.scan_tables import detect_scanned_tables

TEXT_MIN_CHARS = 20


def _clean_text(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _words_to_lines(words: Iterable[dict]) -> list[dict]:
    lines: list[dict] = []
    for word in sorted(words, key=lambda w: (round(w["top"], 1), w["x0"])):
        if not word.get("text"):
            continue
        top = word["top"]
        current = lines[-1] if lines else None
        if current and abs(current["top"] - top) <= 2.5:
            current["words"].append(word)
            current["x0"] = min(current["x0"], word["x0"])
            current["x1"] = max(current["x1"], word["x1"])
            current["bottom"] = max(current["bottom"], word["bottom"])
        else:
            lines.append(
                {
                    "top": top,
                    "x0": word["x0"],
                    "x1": word["x1"],
                    "bottom": word["bottom"],
                    "words": [word],
                }
            )
    for line in lines:
        line["text"] = " ".join(word["text"] for word in line["words"])
        offset = 0
        line["word_spans"] = []
        for word in line["words"]:
            line["word_spans"].append(
                {
                    "start": offset,
                    "end": offset + len(word["text"]),
                    "bbox": [word["x0"], word["top"], word["x1"], word["bottom"]],
                }
            )
            offset += len(word["text"]) + 1
    return lines


def _render_pdf_page(pdf: pdfium.PdfDocument, page_index: int, output: Path) -> Path:
    settings = get_settings()
    output.parent.mkdir(parents=True, exist_ok=True)
    page = pdf[page_index]
    bitmap = page.render(scale=settings.render_scale)
    image = bitmap.to_pil()
    image.save(output, format="PNG")
    return output


def _parse_pdf(
    path: Path, render_dir: Path | None, *, y_tolerance: float = 3,
) -> ParsedDocument:
    pages: list[ParsedPage] = []
    pdfium_doc = pdfium.PdfDocument(path)
    try:
        with pdfplumber.open(path) as pdf:
            max_pages = get_settings().max_pdf_pages
            if len(pdf.pages) > max_pages:
                raise DocumentNeedsVision(f"PDF 页数超过上限 {max_pages}")
            for index, raw_page in enumerate(pdf.pages):
                text = _clean_text(raw_page.extract_text(y_tolerance=y_tolerance) or "")
                lines = _words_to_lines(raw_page.extract_words(y_tolerance=y_tolerance))
                blocks: list[ParsedBlock] = []
                order = 0
                for line in lines:
                    blocks.append(
                        ParsedBlock(
                            kind="line",
                            text=line["text"],
                            bbox=(line["x0"], line["top"], line["x1"], line["bottom"]),
                            order=order,
                            metadata={"word_spans": line["word_spans"]},
                        )
                    )
                    order += 1

                tables = raw_page.find_tables()
                for table_index, table in enumerate(tables):
                    extracted = table.extract(y_tolerance=y_tolerance)
                    if not extracted:
                        continue
                    table_rows = table.rows
                    x_edges = sorted({x for cell in table.cells for x in (cell[0], cell[2])})
                    y_edges = sorted({y for cell in table.cells for y in (cell[1], cell[3])})
                    cells: list[ParsedCell] = []
                    for row_index, row in enumerate(extracted):
                        for col_index, value in enumerate(row):
                            value = _clean_text(value or "")
                            # None denotes a merged-cell continuation; "" is a real empty cell.
                            bbox = table_rows[row_index].cells[col_index]
                            if bbox is None:
                                continue
                            cells.append(
                                ParsedCell(
                                    row=row_index,
                                    col=col_index,
                                    row_span=y_edges.index(bbox[3]) - y_edges.index(bbox[1]),
                                    col_span=x_edges.index(bbox[2]) - x_edges.index(bbox[0]),
                                    text=value,
                                    bbox=tuple(float(v) for v in bbox),
                                    native_locator={
                                        "table_index": table_index,
                                        "row": row_index,
                                        "col": col_index,
                                        "coordinate_quality": "exact",
                                    },
                                )
                            )
                    blocks.append(
                        ParsedBlock(
                            kind="table",
                            text="\n".join(
                                " | ".join(_clean_text(c or "") for c in row) for row in extracted
                            ),
                            bbox=tuple(float(v) for v in table.bbox),
                            order=order,
                            cells=tuple(cells),
                            metadata={"table_index": table_index},
                        )
                    )
                    if extracted:
                        table_top = float(table.bbox[1]) - 2
                        table_bottom = float(table.bbox[3]) + 2
                        line_blocks = [block for block in blocks if block.kind == "line"]
                        non_line_blocks = [block for block in blocks if block.kind != "line"]
                        line_blocks = [
                            block
                            for block in line_blocks
                            if not (
                                block.bbox[1] >= table_top - 2
                                and block.bbox[3] <= table_bottom + 2
                                and block.bbox[0] >= table.bbox[0] - 2
                                and block.bbox[2] <= table.bbox[2] + 2
                            )
                        ]
                        blocks = sorted(
                            [*non_line_blocks, *line_blocks],
                            key=lambda block: (block.bbox[1], block.bbox[0]),
                        )
                        for order_index, block in enumerate(blocks):
                            blocks[order_index] = ParsedBlock(
                                kind=block.kind,
                                text=block.text,
                                bbox=block.bbox,
                                order=order_index,
                                cells=block.cells,
                                metadata=block.metadata,
                            )
                        order = len(blocks)
                    order += 1

                render_path = None
                if render_dir is not None:
                    render_path = _render_pdf_page(
                        pdfium_doc, index, render_dir / f"page-{index + 1:04d}.png"
                    )
                flags = []
                if render_path and raw_page.images and not any(
                    block.kind == "table" for block in blocks
                ):
                    scan_tables = detect_scanned_tables(
                        render_path, float(raw_page.width), float(raw_page.height),
                    )
                    for table in scan_tables:
                        blocks.append(ParsedBlock(
                            kind=table.kind, text=table.text, bbox=table.bbox,
                            order=len(blocks), cells=table.cells, metadata=table.metadata,
                        ))
                    if scan_tables:
                        flags.append("scan_table_geometry")
                mode = (
                    "needs_vision" if len(text) < TEXT_MIN_CHARS and raw_page.images else "text"
                )
                if mode == "needs_vision":
                    flags.append("no_text_layer")
                pages.append(
                    ParsedPage(
                        page_index=index,
                        width=float(raw_page.width),
                        height=float(raw_page.height),
                        text=text,
                        blocks=tuple(blocks),
                        parse_mode=mode,
                        quality_flags=tuple(flags),
                        render_path=render_path,
                        native_locator={"kind": "pdf", "page_index": index},
                    )
                )
    finally:
        pdfium_doc.close()
    return ParsedDocument(pages=tuple(pages), metadata={"page_count": len(pages)})


def _parse_office(path: Path, render_dir: Path | None) -> ParsedDocument:
    if path.suffix.lower() in {".docx", ".xlsx"}:
        _validate_zip(path)
    converted = render_office_pdf(path)
    # Font substitution can shift CJK glyph baselines slightly on Office exports.
    parsed = _parse_pdf(converted, render_dir, y_tolerance=4)
    labels: dict[int, str] = {}
    sheets = []
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = [sheet.title for sheet in workbook if sheet.sheet_state == "visible"]
        workbook.close()
        reader = PdfReader(converted)
        for destination in reader.outline:
            if isinstance(destination, list):
                continue
            if destination.title in sheets:
                index = reader.get_destination_page_number(destination)
                if index is not None:
                    labels[index] = destination.title
    pages = []
    label = None
    for page in parsed.pages:
        label = labels.get(page.page_index, label)
        pages.append(ParsedPage(
            page_index=page.page_index, width=page.width, height=page.height,
            text=page.text, blocks=page.blocks, printed_label=label,
            parse_mode=page.parse_mode, quality_flags=page.quality_flags,
            render_path=page.render_path,
            native_locator={
                "kind": path.suffix[1:].lower(), "page_index": page.page_index,
                "rendered_as": "pdf", **({"sheet": label} if label else {}),
            },
        ))
    return ParsedDocument(
        pages=tuple(pages),
        metadata={**parsed.metadata, "preview_engine": "libreoffice", "sheets": sheets},
    )


def _parse_csv(path: Path) -> ParsedDocument:
    raw = path.read_text(encoding="utf-8-sig")
    rows = list(csv.reader(io.StringIO(raw)))
    cells = tuple(
        ParsedCell(
            row=row_index,
            col=col_index,
            row_span=1,
            col_span=1,
            text=value,
            bbox=(col_index, row_index, col_index + 1, row_index + 1),
            native_locator={"row": row_index + 1, "col": col_index + 1},
        )
        for row_index, row in enumerate(rows)
        for col_index, value in enumerate(row)
    )
    block = ParsedBlock(
        kind="table",
        text="\n".join(" | ".join(row) for row in rows),
        bbox=(0, 0, max((len(row) for row in rows), default=1), max(len(rows), 1)),
        order=0,
        cells=cells,
    )
    page = ParsedPage(
        page_index=0,
        width=max((len(row) for row in rows), default=1),
        height=max(len(rows), 1),
        text=block.text,
        blocks=(block,),
        parse_mode="text",
        quality_flags=("spreadsheet_coordinates",),
        native_locator={"kind": "csv"},
    )
    return ParsedDocument(pages=(page,), metadata={"page_count": 1})


def parse_document(path: Path, render_root: Path | None = None) -> ParsedDocument:
    suffix = path.suffix.lower()
    render_dir = render_root / path.stem if render_root else None
    if suffix == ".pdf":
        return _parse_pdf(path, render_dir)
    if suffix in {".doc", ".docx", ".xlsx"}:
        return _parse_office(path, render_dir)
    if suffix == ".csv":
        return _parse_csv(path)
    raise DocumentNeedsVision(f"暂不支持直接解析 {suffix or '未知格式'} 文件")


def _validate_zip(path: Path) -> None:
    settings = get_settings()
    try:
        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > settings.max_zip_entries:
                raise DocumentNeedsVision("Office 文件压缩条目过多")
            if sum(info.file_size for info in infos) > settings.max_zip_uncompressed_bytes:
                raise DocumentNeedsVision("Office 文件解压后过大")
    except zipfile.BadZipFile as exc:
        raise DocumentNeedsVision("Office 文件结构无效") from exc
