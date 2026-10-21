import hashlib
import math
from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageFilter

from app.services.office_preview import office_executable
from app.services.parser import parse_document
from app.services.scan_tables import detect_scanned_tables


def scanned_form(path: Path, angle: float = 0) -> Path:
    image = Image.new("RGB", (1000, 850), "white")
    draw = ImageDraw.Draw(image)
    for y in (200, 280, 440, 520):
        draw.line((100, y, 850, y), fill="black", width=3)
    draw.line((350, 360, 850, 360), fill="black", width=3)
    for x in (100, 600, 850):
        draw.line((x, 200, x, 520), fill="black", width=3)
    draw.line((350, 280, 350, 520), fill="black", width=3)
    draw.text((130, 225), "Equipment details", fill="black", font_size=22)
    draw.text((640, 225), "Quantity", fill="black", font_size=22)
    draw.text((130, 310), "Model", fill="black", font_size=20)
    draw.rectangle((598, 322, 602, 325), fill="white")
    image = image.filter(ImageFilter.GaussianBlur(0.35))
    if angle:
        image = image.rotate(angle, resample=Image.Resampling.BICUBIC, fillcolor="white")
    image.save(path)
    return path


@pytest.mark.parametrize("angle", [0, 2.2, -3])
def test_scanned_table_preserves_merged_cells_and_original_coordinates(tmp_path, angle):
    source = scanned_form(tmp_path / "scan.png", angle)
    tables = detect_scanned_tables(source, 500, 425)
    assert len(tables) == 1
    assert len(tables[0].cells) == 10
    assert any(cell.row_span == 2 for cell in tables[0].cells)
    assert any(cell.col_span == 2 for cell in tables[0].cells)
    assert sum(cell.native_locator["scan_has_content"] for cell in tables[0].cells) == 3
    last = max(tables[0].cells, key=lambda cell: (cell.row, cell.col))
    x, y = 725 - 500, 480 - 425
    expected = (
        (500 + x * math.cos(math.radians(angle)) + y * math.sin(math.radians(angle))) / 2,
        (425 - x * math.sin(math.radians(angle)) + y * math.cos(math.radians(angle))) / 2,
    )
    polygon = last.native_locator["polygon"]
    center = (sum(point[0] for point in polygon) / 4, sum(point[1] for point in polygon) / 4)
    assert math.dist(expected, center) < 3


def test_scanned_template_to_fill_annotations_is_offline(tmp_path, monkeypatch):
    import io

    from fastapi.testclient import TestClient
    from test_offline_workflow import prepare

    from app.main import app

    async def reject_model(*_args, **_kwargs):
        raise AssertionError("Scan parsing attempted a model call")

    monkeypatch.setattr("app.services.handlers.transcribe_scanned_page", reject_model)
    monkeypatch.setattr("httpx.AsyncClient.post", reject_model)
    source = scanned_form(tmp_path / "scan.png", 2.2)
    pdf = source.with_suffix(".pdf")
    with Image.open(source) as image:
        image.save(pdf, "PDF", resolution=144)
    with TestClient(app) as client:
        _, documents, run = prepare(client, pdf, ["tender", "template"])
        pages = client.get(f"/api/v2/documents/{documents[1]['id']}/pages").json()
        assert "scan_table_geometry" in pages[0]["quality_flags"]
        assert pages[0]["parse_mode"] == "needs_vision"
        response = client.post(f"/api/v2/analysis-runs/{run['id']}/fill-plan", json={})
        assert response.status_code == 201, response.text
        plan = response.json()
        assert len(plan["targets"]) == 7  # The three occupied cells must remain untouched.
        assert all(
            annotation["coordinate_quality"] == "scan_lines" and len(annotation["polygon"]) == 4
            for asset in plan["assets"] for annotation in asset["annotations"]
        )
        target = plan["targets"][0]
        renamed = client.patch(
            f"/api/v2/fill-targets/{target['id']}", json={"label": "Equipment model"},
        )
        assert renamed.status_code == 200, renamed.text
        assert renamed.json()["assets"][0]["annotations"][0]["polygon"]
        saved = client.patch(
            f"/api/v2/fill-suggestions/{target['suggestions'][0]['id']}",
            json={"value": "智能控制柜", "status": "edited"},
        )
        assert saved.status_code == 200, saved.text
        asset = plan["assets"][0]
        image = client.get(f"/api/v2/annotation-assets/{asset['id']}/image")
        assert image.status_code == 200
        original = client.get(f"/api/v2/pages/{pages[0]['id']}/image")
        with Image.open(io.BytesIO(image.content)) as annotated, Image.open(
            io.BytesIO(original.content)
        ) as original_page:
            # Header text is outside every fill polygon.
            assert annotated.crop((130, 225, 330, 250)).tobytes() == original_page.crop(
                (130, 225, 330, 250)
            ).tobytes()
        moved = client.patch(
            f"/api/v2/fill-targets/{target['id']}", json={"bbox": [180, 180, 280, 210]},
        )
        assert moved.status_code == 200, moved.text
        annotation = next(
            item for item in moved.json()["assets"][0]["annotations"]
            if item["target_id"] == target["id"]
        )
        assert annotation["coordinate_quality"] == "manual"
        assert annotation["polygon"] is None


def test_plain_scan_text_is_not_detected_as_a_table(tmp_path):
    image = Image.new("RGB", (1000, 850), "white")
    draw = ImageDraw.Draw(image)
    for y in range(100, 600, 40):
        draw.text((100, y), "Tender requirements and supporting evidence", font_size=20)
    source = tmp_path / "text.png"
    image.save(source)
    assert detect_scanned_tables(source, 500, 425) == []


@pytest.mark.skipif(office_executable() is None, reason="LibreOffice not installed")
def test_word_print_pages_are_rendered_without_modifying_source(tmp_path):
    from docx import Document
    from docx.shared import Cm

    document = Document()
    document.add_heading("投标响应文件", 0)
    document.add_paragraph("第一部分：企业及设备信息。")
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("设备名称", "规格型号", "数量"), strict=True):
        cell.text = value
    document.add_page_break()
    document.add_heading("交付与服务承诺", 1)
    document.add_paragraph("合同签订后 30 日内交付，提供现场调试支持。")
    document.sections[0].left_margin = Cm(2)
    path = tmp_path / "response.docx"
    document.save(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    parsed = parse_document(path, tmp_path / "rendered")
    assert len(parsed.pages) == 2
    assert all(page.render_path and page.render_path.exists() for page in parsed.pages)
    assert "交付与服务承诺" in parsed.pages[1].text
    assert all(page.native_locator["rendered_as"] == "pdf" for page in parsed.pages)
    assert any(
        cell.text == "" and cell.bbox[2] - cell.bbox[0] < parsed.pages[0].width
        for block in parsed.pages[0].blocks for cell in block.cells
    )
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


@pytest.mark.skipif(office_executable() is None, reason="LibreOffice not installed")
def test_excel_preview_respects_sheets_print_areas_and_orientation(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import Border, Font, Side

    workbook = Workbook()
    for sheet, title in (
        (workbook.active, "报价明细"), (workbook.create_sheet(), "交付安排"),
    ):
        sheet.title = title
        sheet.append(["项目", "数量", "单价"])
        sheet.append(["控制柜", 2, 1000])
        sheet.append(["配电柜", 3, 2000])
        sheet["A5"] = "合计数量"
        sheet["B5"] = "=SUM(B2:B3)"
        sheet["H30"] = "OUTSIDE_PRINT_AREA"
        for column in ("A", "B", "C"):
            sheet.column_dimensions[column].width = 24
        for row in sheet.iter_rows(min_row=1, max_row=5, max_col=3):
            for cell in row:
                cell.font = Font(name="Noto Sans CJK SC", size=12)
                cell.border = Border(**{
                    side: Side(style="thin", color="000000")
                    for side in ("left", "right", "top", "bottom")
                })
        sheet.print_area = "A1:C5"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
    hidden = workbook.create_sheet("内部草稿")
    hidden["A1"] = "HIDDEN_SHEET_CONTENT"
    hidden.sheet_state = "hidden"
    path = tmp_path / "pricing.xlsx"
    workbook.save(path)
    before = path.read_bytes()
    parsed = parse_document(path, tmp_path / "rendered")
    assert len(parsed.pages) == 2
    assert all(page.width > page.height for page in parsed.pages)
    assert all(page.render_path and page.render_path.exists() for page in parsed.pages)
    text = "\n".join(page.text for page in parsed.pages)
    assert "控制柜" in text
    assert "OUTSIDE_PRINT_AREA" not in text
    assert "HIDDEN_SHEET_CONTENT" not in text
    assert "5" in text
    assert [page.printed_label for page in parsed.pages] == ["报价明细", "交付安排"]
    assert all(page.native_locator["kind"] == "xlsx" for page in parsed.pages)
    assert path.read_bytes() == before
